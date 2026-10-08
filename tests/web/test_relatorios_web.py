"""Fase 3: Relatórios (quatro abas, período na URL e exportação em CSV e Excel)."""

from datetime import date
from io import BytesIO

from openpyxl import load_workbook

from tests.web.conftest import entrar

HX = {"HX-Request": "true", "HX-Target": "painel-aba"}


def test_resultado_por_moto_mostra_valores_resumo_e_barras(cliente, base_relatorios):
    entrar(cliente)
    html = cliente.get("/relatorios").text
    assert "BRA-2E19" in html and "QRS-4T21" in html and "R$ 1.000,00" in html and "R$ 850,00" in html
    assert "Resultado total: <strong>R$ 750,00</strong> em 2 motos." in html
    assert "Melhor: <strong>BRA-2E19</strong> (R$ 850,00). Pior: <strong>QRS-4T21</strong> (R$ -100,00)." in html
    assert html.index("BRA-2E19") < html.index("QRS-4T21")  # maior resultado primeiro
    assert 'class="barra-rel-ok" width="100"' in html and 'class="barra-rel-perigo"' not in html  # negativo fica sem barra
    assert "Cauções não compõem receita" in html and 'aria-current' in html
    assert "style=" not in html.split("<main", 1)[1]  # a CSP não aceita estilo inline


def test_periodo_padrao_e_o_informado(cliente, base_relatorios):
    entrar(cliente)
    html = cliente.get("/relatorios").text
    assert base_relatorios.consultas[-1] == (base_relatorios.hoje.replace(day=1), base_relatorios.hoje)
    assert f'value="{base_relatorios.hoje.replace(day=1).isoformat()}"' in html
    html = cliente.get("/relatorios?de=2026-01-05&ate=2026-03-31").text
    assert base_relatorios.consultas[-1] == (date(2026, 1, 5), date(2026, 3, 31))
    assert "05/01/2026 a 31/03/2026" in html and 'value="2026-01-05"' in html and 'value="2026-03-31"' in html


def test_periodo_invertido_avisa_e_nao_calcula(cliente, base_relatorios):
    entrar(cliente)
    html = cliente.get("/relatorios?de=2026-10-08&ate=2026-10-01").text
    assert "A data final deve ser igual ou posterior à inicial." in html and base_relatorios.consultas == []
    assert "Exportar CSV" not in html and "BRA-2E19" not in html


def test_abas_viajam_pelo_painel_com_o_periodo(cliente, base_relatorios):
    entrar(cliente)
    trecho = cliente.get("/relatorios/abas/fluxo?de=2026-01-01&ate=2026-10-08", headers=HX).text
    assert "<html" not in trecho and 'id="painel-aba"' in trecho and 'hx-swap-oob="true"' in trecho
    assert "Setembro de 2026" in trecho and "(parcial)" in trecho and "R$ 800,00" in trecho
    assert 'hx-push-url="/relatorios?aba=custo&amp;de=2026-01-01&amp;ate=2026-10-08"' in trecho
    pagina = cliente.get("/relatorios?aba=fluxo", headers=HX).text
    assert "<html" not in pagina
    assert "<html" in cliente.get("/relatorios?aba=fluxo").text
    assert cliente.get("/relatorios/abas/inexistente").status_code == 404
    assert "Resultado por moto" in cliente.get("/relatorios?aba=invalida").text


def test_custo_de_manutencao_por_modelo_e_por_moto(cliente, base_relatorios):
    entrar(cliente)
    modelo = cliente.get("/relatorios?aba=custo").text
    assert "Custo de manutenção por modelo" in modelo and "Factor &lt;i&gt;150&lt;/i&gt;" in modelo and "<i>150</i>" not in modelo
    assert "R$ 300,00" in modelo and "Custo total: <strong>R$ 400,00</strong> em 2 modelos." in modelo
    assert modelo.index("Factor") < modelo.index("CG 160")  # maior custo primeiro
    moto = cliente.get("/relatorios?aba=custo&visao=moto").text
    assert "Custo de manutenção por moto" in moto and "500 km" in moto and "R$ 0,20" in moto
    assert moto.index("QRS-4T21") < moto.index("BRA-2E19") and "—" in moto
    assert "Custo de manutenção por modelo" in cliente.get("/relatorios?aba=custo&visao=x").text


def test_inadimplencia_ignora_o_periodo_e_mostra_os_indicadores(cliente, base_relatorios):
    entrar(cliente)
    html = cliente.get("/relatorios?aba=inadimplencia&de=2026-01-01&ate=2026-01-02").text
    assert "R$ 380,00" in html and "13,6%" in html and "Maria &lt;b&gt;Silva&lt;/b&gt;" in html and "37 dia(s)" in html
    assert "R$ 554,00" in html and "BRA-2E19" in html and 'name="de"' not in html
    assert base_relatorios.consultas == []  # a posição de hoje não consulta o período
    assert "Posição de hoje" in html and "01/01/2026 a 02/01/2026" not in html
    base_relatorios.inadimplencia = {"linhas": [], "total_atraso": 0, "clientes": 0, "percentual_carteira": None}
    vazio = cliente.get("/relatorios?aba=inadimplencia").text
    assert "Nenhuma cobrança em atraso" in vazio and "—" in vazio and "Exportar" not in vazio


def test_sem_dados_no_periodo_mostra_estado_vazio_e_esconde_exportacao(cliente, base_relatorios):
    entrar(cliente)
    base_relatorios.fluxo = []
    html = cliente.get("/relatorios?aba=fluxo").text
    assert "Nada neste período" in html and "Exportar" not in html


def test_exportar_csv_traz_os_dados_da_tabela(cliente, base_relatorios):
    entrar(cliente)
    resposta = cliente.get("/relatorios/exportar?aba=resultado&de=2026-10-01&ate=2026-10-08&formato=csv")
    assert resposta.status_code == 200 and resposta.headers["content-type"].startswith("text/csv")
    assert resposta.headers["content-disposition"] == 'attachment; filename="relatorio_resultado_por_moto.csv"'
    texto = resposta.content.decode("utf-8-sig")
    linhas = texto.splitlines()
    assert linhas[0] == "Placa;Modelo;Receita recebida;Manutenção;Documentos;Resultado;Km rodados;Custo por km"
    assert linhas[1].startswith("BRA-2E19;CG 160;1000;100;50;850;500;0,20") and linhas[2].startswith("QRS-4T21;")
    assert base_relatorios.consultas[-1] == (date(2026, 10, 1), date(2026, 10, 8))


def test_exportar_excel_e_nomes_de_arquivo_de_cada_aba(cliente, base_relatorios):
    entrar(cliente)
    resposta = cliente.get("/relatorios/exportar?aba=fluxo&formato=xlsx")
    assert resposta.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert resposta.headers["content-disposition"].endswith('relatorio_fluxo_de_caixa.xlsx"')
    aba = load_workbook(BytesIO(resposta.content)).active
    assert [c.value for c in aba[1]] == ["Mês", "Recebido", "Manutenção", "Documentos", "Líquido"] and aba.max_row == 3
    nomes = {
        "custo": "relatorio_custo_manutencao_modelo.csv", "inadimplencia": "relatorio_inadimplencia.csv",
    }
    for chave, arquivo in nomes.items():
        assert cliente.get(f"/relatorios/exportar?aba={chave}").headers["content-disposition"].endswith(arquivo + '"')
    assert cliente.get("/relatorios/exportar?aba=custo&visao=moto").headers["content-disposition"].endswith('moto.csv"')


def test_exportar_protege_contra_formula_no_csv(cliente, base_relatorios):
    entrar(cliente)
    base_relatorios.resultado[0]["modelo"] = "=HYPERLINK(\"http://x\")"
    texto = cliente.get("/relatorios/exportar?aba=resultado").content.decode("utf-8-sig")
    assert "'=HYPERLINK" in texto


def test_exportar_recusa_formato_periodo_invertido_e_sem_dados(cliente, base_relatorios):
    entrar(cliente)
    assert cliente.get("/relatorios/exportar?aba=resultado&formato=pdf").status_code == 404
    assert cliente.get("/relatorios/exportar?aba=resultado&de=2026-10-08&ate=2026-10-01").status_code == 422
    base_relatorios.fluxo = []
    assert cliente.get("/relatorios/exportar?aba=fluxo").status_code == 404


def test_relatorios_exigem_dono(cliente, base_relatorios):
    assert cliente.get("/relatorios").status_code == 303
    assert cliente.get("/relatorios/exportar").status_code == 303
    entrar(cliente, identificador="123.456.789-09", senha="senha-locatario")
    assert cliente.get("/relatorios").status_code == 403
    assert cliente.get("/relatorios/exportar").status_code == 403
    assert cliente.get("/relatorios/abas/fluxo").status_code == 403
